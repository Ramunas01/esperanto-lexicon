#!/usr/bin/env python3
"""Concept dedup — one concept per Esperanto word (decided 2026-10-14).

Concept identity is the Esperanto word ("option (a)"): English synonyms that
Esperanto joins (``achieve``/``accomplish`` → ``atingi``) are ONE concept with
several English rows, not several concepts. This tool has three subcommands:

``worksheet``
    Read-only. Classifies every ``eo_word`` carried by more than one concept,
    and every inflected EN row that sits on a different concept than its lemma,
    and writes a review workbook. Nothing in the DB changes.

``merge-synonyms``
    Merges each SAFE group — every member has at least one EN word that is a
    sense of the ``eo_word``'s ESPDIC gloss (a true synonym, not a wrong
    anchor), and the members agree on all Esperanto identity fields. The
    survivor keeps its id; the others' ``concept_lang`` rows are re-pointed to
    it (an identical duplicate row is dropped), their ``concept_root`` and
    ``concept`` rows are removed, and a reversible snapshot is written to
    ``concept_merge``. Suspect groups are left for human review.

``fold-inflections``
    Moves EN ``concept_lang`` rows that are inflected surface forms
    (``methodologies``, ``achieving``) of another EN word ON THE SAME concept
    into ``inflected_forms`` (``form_description='concept_row_fold'``), keeping
    the row's own tier. Only rows whose tier is >= the lemma's tier are moved,
    so no surface form resolves to a different tier afterwards.

Every write subcommand is ``--dry-run`` by default; ``--commit`` takes a
timestamped backup, runs in one transaction, and commits only if the post-write
audit passes (the (lang, word, tier) metric invariant, no orphans, the
``eo_root == concept_root head`` invariant, and the subcommand's own counts).
No tier of any existing word is changed.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sqlite3
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Callable

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from lexicon.schema import create_concept_merge_schema  # noqa: E402

MERGE_SOURCE = "concept_dedup_v1"
FOLD_DESCRIPTION = "concept_row_fold"

# Penn tags of inflected (non-base) forms; a row whose word carries one of these
# and whose lemma is another EN word on the same concept is a fold candidate.
INFLECTED_TAGS = ("NNS", "VBD", "VBN", "VBG", "VBZ", "JJR", "JJS", "RBR", "RBS")

# The Esperanto identity of a concept: members of one group must not carry two
# different real values for any of these (eo_word agrees by construction).
IDENTITY_FIELDS = ("eo_root", "eo_pos", "eo_prefix", "eo_suffix", "eo_status")

# Values meaning "unknown/empty": older writers stored '' for no affix, the
# gap-fill writers NULL, and 'X' is the UD tag for an unknown part of speech.
# Empty is compatible with anything; on merge the survivor's empty fields are
# filled from the other members.
EMPTY_VALUES = (None, "", "X")

# (word) -> (lemma, penn_tag); injected so tests need no spaCy model.
Lemmatizer = Callable[[str], tuple[str, str]]


def spacy_lemmatizer() -> Lemmatizer:
    """Return a cached single-word lemmatizer backed by ``en_core_web_sm``."""
    import spacy

    nlp = spacy.load("en_core_web_sm", disable=["parser", "ner"])
    cache: dict[str, tuple[str, str]] = {}

    def lemmatize(word: str) -> tuple[str, str]:
        if word not in cache:
            doc = nlp(word)
            if len(doc) == 1:
                cache[word] = (doc[0].lemma_.lower(), doc[0].tag_)
            else:
                cache[word] = (word, "")
        return cache[word]

    return lemmatize


# --------------------------------------------------------------------------- #
# ESPDIC glosses
# --------------------------------------------------------------------------- #
def load_espdic_senses(espdic: Path) -> dict[str, set[str]]:
    """Return ``{eo_word: {english sense, ...}}`` from ``espdic.txt``.

    Senses are the comma/semicolon-separated gloss items, lowercased, with
    parentheticals and a leading ``to``/``a``/``an``/``the`` removed.
    """
    senses: dict[str, set[str]] = defaultdict(set)
    if not espdic.exists():
        return senses
    for line in espdic.read_text(encoding="utf-8").splitlines():
        if " : " not in line:
            continue
        eo, gloss = line.split(" : ", 1)
        gloss = re.sub(r"\([^)]*\)", "", gloss.lower())
        for part in re.split(r"[,;]", gloss):
            sense = re.sub(r"^(to|a|an|the)\s+", "", part.strip())
            if sense:
                senses[eo.strip().lower()].add(sense)
    return senses


# --------------------------------------------------------------------------- #
# Group classification
# --------------------------------------------------------------------------- #
@dataclass
class Member:
    """One concept in an ``eo_word`` group."""

    cid: int
    identity: tuple
    synset: str | None
    en_words: list[tuple[str, str, int | None, str]]  # (word, pos, tier, source)
    in_gloss: bool = False


@dataclass
class Group:
    """All concepts anchored to one ``eo_word``, with the merge verdict."""

    eo_word: str
    members: list[Member]
    gloss: set[str] = field(default_factory=set)
    verdict: str = ""  # safe_synonyms | suspect_anchor | no_espdic | identity_conflict
    survivor: int | None = None


def load_groups(conn: sqlite3.Connection) -> list[Group]:
    """Return every ``eo_word`` carried by more than one concept."""
    by_word: dict[str, list[Member]] = defaultdict(list)
    cols = ", ".join(IDENTITY_FIELDS)
    for row in conn.execute(
        f"SELECT id, eo_word, wordnet_synset, {cols} FROM concept "
        "WHERE eo_word IS NOT NULL ORDER BY id"
    ):
        cid, eo_word, synset, *identity = row
        en = conn.execute(
            "SELECT LOWER(word), pos, tier, source FROM concept_lang "
            "WHERE concept_id = ? AND lang = 'en' ORDER BY tier, word",
            (cid,),
        ).fetchall()
        by_word[eo_word].append(Member(cid, tuple(identity), synset, en))
    return [Group(w, ms) for w, ms in sorted(by_word.items()) if len(ms) > 1]


def _root_set(conn: sqlite3.Connection, cid: int) -> tuple:
    return tuple(
        conn.execute(
            "SELECT position, root, is_head FROM concept_root "
            "WHERE concept_id = ? ORDER BY position",
            (cid,),
        ).fetchall()
    )


def merged_identity(members: list[Member]) -> dict[str, object] | None:
    """Return the merged identity fields, or None if two real values conflict.

    Each field takes its single non-empty value across the members (or stays
    empty if every member is empty). Two different non-empty values = conflict.
    """
    merged: dict[str, object] = {}
    for i, name in enumerate(IDENTITY_FIELDS):
        real = {m.identity[i] for m in members} - set(EMPTY_VALUES)
        if len(real) > 1:
            return None
        merged[name] = real.pop() if real else None
    return merged


def merged_roots(conn: sqlite3.Connection, members: list[Member]) -> tuple | None:
    """Return the members' single non-empty ``concept_root`` set, or None on conflict."""
    real = {_root_set(conn, m.cid) for m in members} - {()}
    if len(real) > 1:
        return None
    return real.pop() if real else ()


def pick_survivor(members: list[Member]) -> int:
    """Survivor: a member with a WordNet synset, then lowest tier, then lowest id."""

    def key(m: Member) -> tuple:
        tiers = [t for _, _, t, _ in m.en_words if t is not None]
        return (m.synset is None, min(tiers) if tiers else 99, m.cid)

    return min(members, key=key).cid


def classify_groups(
    conn: sqlite3.Connection,
    groups: list[Group],
    senses: dict[str, set[str]],
    lemmatize: Lemmatizer,
) -> None:
    """Set ``verdict`` (and ``survivor`` for safe groups) on every group in place."""
    for g in groups:
        g.gloss = senses.get(g.eo_word.lower(), set())
        for m in g.members:
            m.in_gloss = any(
                w in g.gloss or lemmatize(w)[0] in g.gloss for w, _, _, _ in m.en_words
            )
        lifecycle = _has_table(conn, "concept_lifecycle") and conn.execute(
            "SELECT 1 FROM concept_lifecycle WHERE concept_id IN (%s)"
            % ",".join("?" * len(g.members)),
            tuple(m.cid for m in g.members),
        ).fetchone()
        conflict = (
            merged_identity(g.members) is None
            or merged_roots(conn, g.members) is None
        )
        if conflict or lifecycle:
            g.verdict = "identity_conflict"
        elif not g.gloss:
            g.verdict = "no_espdic"
        elif all(m.in_gloss for m in g.members):
            g.verdict = "safe_synonyms"
            g.survivor = pick_survivor(g.members)
        else:
            g.verdict = "suspect_anchor"


def _has_table(conn: sqlite3.Connection, name: str) -> bool:
    return (
        conn.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (name,)
        ).fetchone()
        is not None
    )


# --------------------------------------------------------------------------- #
# Inflection candidates
# --------------------------------------------------------------------------- #
@dataclass
class InflectedRow:
    """An EN concept_lang row whose word is an inflected form of another EN word."""

    row_id: int
    cid: int
    word: str
    lemma: str
    tag: str
    tier: int | None
    source: str
    lemma_cids: set[int]
    lemma_tier: int | None  # lowest tier of the lemma on the same concept


def find_inflected_rows(
    conn: sqlite3.Connection, lemmatize: Lemmatizer
) -> list[InflectedRow]:
    """Return single-word EN rows that are inflections of another EN lexicon word."""
    where: dict[str, dict[int, list[int | None]]] = defaultdict(lambda: defaultdict(list))
    rows = conn.execute(
        "SELECT id, concept_id, LOWER(word), tier, source FROM concept_lang "
        "WHERE lang = 'en' AND word NOT LIKE '% %'"
    ).fetchall()
    for _, cid, w, tier, _ in rows:
        where[w][cid].append(tier)
    found = []
    for row_id, cid, w, tier, source in rows:
        lemma, tag = lemmatize(w)
        if lemma == w or tag not in INFLECTED_TAGS or lemma not in where:
            continue
        same = [t for t in where[lemma].get(cid, []) if t is not None]
        found.append(
            InflectedRow(
                row_id, cid, w, lemma, tag, tier, source,
                set(where[lemma]), min(same) if same else None,
            )
        )
    return found


def foldable(r: InflectedRow) -> bool:
    """True if the row can move to ``inflected_forms`` without any tier change."""
    return (
        r.cid in r.lemma_cids
        and r.tier is not None
        and r.lemma_tier is not None
        and r.tier >= r.lemma_tier
    )


# --------------------------------------------------------------------------- #
# Audits
# --------------------------------------------------------------------------- #
def word_tier_triples(conn: sqlite3.Connection) -> set[tuple[str, str, int | None]]:
    """The metric-relevant content: every (lang, lowercased word, tier)."""
    return set(
        conn.execute("SELECT DISTINCT lang, LOWER(word), tier FROM concept_lang")
    )


def structural_audit(conn: sqlite3.Connection) -> dict[str, int]:
    """Orphan rows and ``eo_root == concept_root head`` mismatches (all should be 0)."""
    q = lambda sql: conn.execute(sql).fetchone()[0]  # noqa: E731
    return {
        "orphan_concept_lang": q(
            "SELECT COUNT(*) FROM concept_lang WHERE concept_id NOT IN "
            "(SELECT id FROM concept)"
        ),
        "orphan_concept_root": q(
            "SELECT COUNT(*) FROM concept_root WHERE concept_id NOT IN "
            "(SELECT id FROM concept)"
        ),
        "head_mismatches": q(
            "SELECT COUNT(*) FROM concept c JOIN concept_root cr "
            "ON cr.concept_id = c.id AND cr.is_head = 1 WHERE c.eo_root <> cr.root"
        ),
    }


# --------------------------------------------------------------------------- #
# Writers (callers own the transaction)
# --------------------------------------------------------------------------- #
def merge_group(conn: sqlite3.Connection, g: Group, asof: str) -> dict[str, int]:
    """Merge one safe group into its survivor. Returns moved/dropped counts.

    The survivor's empty identity fields and missing ``concept_root`` rows are
    filled from the other members first; its pre-merge values go into the
    first retired member's snapshot (``survivor_before``) for reversal.
    """
    moved = dropped = 0
    identity = merged_identity(g.members)
    roots = merged_roots(conn, g.members)
    if identity is None or roots is None:
        raise ValueError(f"{g.eo_word}: identity conflict, not mergeable")
    cols = ", ".join(IDENTITY_FIELDS)
    survivor_before = {
        "identity": dict(
            zip(
                IDENTITY_FIELDS,
                conn.execute(
                    f"SELECT {cols} FROM concept WHERE id = ?", (g.survivor,)
                ).fetchone(),
            )
        ),
        "concept_root": [list(r) for r in _root_set(conn, g.survivor)],
    }
    for name, value in identity.items():
        conn.execute(
            f"UPDATE concept SET {name} = ? WHERE id = ? "
            f"AND ({name} IS NULL OR {name} IN ('', 'X'))",
            (value, g.survivor),
        )
    if roots and not survivor_before["concept_root"]:
        donor = next(m.cid for m in g.members if _root_set(conn, m.cid) == roots)
        conn.execute(
            "INSERT INTO concept_root (concept_id, root, position, is_head, tier) "
            "SELECT ?, root, position, is_head, tier FROM concept_root "
            "WHERE concept_id = ?",
            (g.survivor, donor),
        )
    for m in g.members:
        if m.cid == g.survivor:
            continue
        concept_row = conn.execute("SELECT * FROM concept WHERE id = ?", (m.cid,))
        names = [d[0] for d in concept_row.description]
        snapshot: dict = {
            "concept": dict(zip(names, concept_row.fetchone())),
            "concept_root": [
                list(r)
                for r in conn.execute(
                    "SELECT root, position, is_head, tier FROM concept_root "
                    "WHERE concept_id = ? ORDER BY position",
                    (m.cid,),
                )
            ],
            "moved_concept_lang_ids": [],
            "dropped_concept_lang": [],
        }
        if survivor_before is not None:
            snapshot["survivor_before"] = survivor_before
            survivor_before = None
        for row_id, lang, word, pos, cefr, tier, source in conn.execute(
            "SELECT id, lang, word, pos, cefr_level, tier, source FROM concept_lang "
            "WHERE concept_id = ?",
            (m.cid,),
        ).fetchall():
            clash = conn.execute(
                "SELECT cefr_level, tier FROM concept_lang "
                "WHERE concept_id = ? AND lang = ? AND word = ? AND pos = ?",
                (g.survivor, lang, word, pos),
            ).fetchone()
            if clash is None:
                conn.execute(
                    "UPDATE concept_lang SET concept_id = ? WHERE id = ?",
                    (g.survivor, row_id),
                )
                snapshot["moved_concept_lang_ids"].append(row_id)
                moved += 1
            elif clash[1] == tier:
                conn.execute("DELETE FROM concept_lang WHERE id = ?", (row_id,))
                snapshot["dropped_concept_lang"].append(
                    [row_id, lang, word, pos, cefr, tier, source]
                )
                dropped += 1
            else:
                raise ValueError(
                    f"{g.eo_word}: {lang}/{word}/{pos} tier {tier} on concept "
                    f"{m.cid} vs {clash[1]} on survivor {g.survivor}"
                )
        conn.execute("DELETE FROM concept_root WHERE concept_id = ?", (m.cid,))
        conn.execute("DELETE FROM concept WHERE id = ?", (m.cid,))
        conn.execute(
            "INSERT INTO concept_merge "
            "(retired_id, survivor_id, eo_word, source, asof, snapshot) "
            "VALUES (?, ?, ?, ?, ?, ?)",
            (m.cid, g.survivor, g.eo_word, MERGE_SOURCE, asof, json.dumps(snapshot)),
        )
    return {"moved": moved, "dropped": dropped}


def fold_row(conn: sqlite3.Connection, r: InflectedRow) -> bool:
    """Move one foldable row to ``inflected_forms``. Returns True if inserted."""
    before = conn.total_changes
    conn.execute(
        "INSERT OR IGNORE INTO inflected_forms "
        "(inflected_word, lemma, lang, form_description, tier) "
        "VALUES (?, ?, 'en', ?, ?)",
        (r.word, r.lemma, FOLD_DESCRIPTION, r.tier),
    )
    inserted = conn.total_changes > before
    conn.execute("DELETE FROM concept_lang WHERE id = ?", (r.row_id,))
    return inserted


def resolved_tier(conn: sqlite3.Connection, word: str, lemma: str) -> int | None:
    """Best (lowest) tier a surface form reaches by text or lemma — as the analyzer does."""
    row = conn.execute(
        "SELECT MIN(tier) FROM concept_lang WHERE lang = 'en' AND LOWER(word) IN (?, ?)",
        (word, lemma),
    ).fetchone()
    return row[0]


# --------------------------------------------------------------------------- #
# Worksheet
# --------------------------------------------------------------------------- #
def write_worksheet(
    groups: list[Group], inflected: list[InflectedRow], out: Path
) -> None:
    """Write the review workbook: one sheet for groups, one for cross-concept inflections."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "eo_word_groups"
    ws.append(
        ["eo_word", "verdict", "espdic_gloss", "concept_id", "survivor", "in_gloss",
         "en_words (word/pos/tier/source)", "decision", "note"]
    )
    order = {"suspect_anchor": 0, "no_espdic": 1, "identity_conflict": 2,
             "safe_synonyms": 3}
    for g in sorted(groups, key=lambda g: (order[g.verdict], g.eo_word)):
        for m in g.members:
            ws.append([
                g.eo_word, g.verdict, ", ".join(sorted(g.gloss)), m.cid,
                "yes" if m.cid == g.survivor else "",
                "yes" if m.in_gloss else "NO",
                "; ".join(f"{w}/{p}/T{t}/{s}" for w, p, t, s in m.en_words),
                "", "",
            ])
    ws2 = wb.create_sheet("cross_concept_inflections")
    ws2.append(["word", "lemma", "tag", "tier", "source", "concept_id",
                "lemma_concept_ids", "decision", "note"])
    for r in sorted(inflected, key=lambda r: r.word):
        if r.cid in r.lemma_cids:
            continue
        ws2.append([r.word, r.lemma, r.tag, r.tier, r.source, r.cid,
                    ", ".join(map(str, sorted(r.lemma_cids))), "", ""])
    readme = wb.create_sheet("README", 0)
    for line in (
        "Concept dedup review — concept identity = the Esperanto word (2026-10-14).",
        "",
        "Sheet eo_word_groups: concepts sharing one eo_word.",
        "  safe_synonyms: every member's EN word is in the ESPDIC gloss — merged",
        "    automatically by merge-synonyms (listed for reference, no action).",
        "  suspect_anchor: some member (in_gloss=NO) is not in the gloss — it may be",
        "    a looser synonym (merge), an inflection, or a WRONG ANCHOR.",
        "  no_espdic: the eo_word is not in ESPDIC at all.",
        "  identity_conflict: members disagree on eo_root/eo_pos/affixes/roots.",
        "  decision per member row: merge | keep_separate | reanchor:<eo_word> | drop",
        "",
        "Sheet cross_concept_inflections: an inflected EN word on a different",
        "  concept than its lemma (achieved on c493, achieve on c20).",
        "  decision: fold (-> inflected_forms of the lemma) | keep (separate lexeme)",
    ):
        readme.append([line])
    for sheet in (ws, ws2):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def _backup(lexicon: Path, tag: str) -> Path:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = lexicon.with_suffix(f".db.bak-{tag}-{stamp}")
    shutil.copy2(lexicon, backup)
    return backup


def _counts(conn: sqlite3.Connection) -> dict[str, int]:
    return {
        t: conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
        for t in ("concept", "concept_lang", "concept_root", "inflected_forms")
    }


def cmd_worksheet(conn: sqlite3.Connection, args: argparse.Namespace) -> int:
    """Classify groups and inflections; write the review workbook."""
    lem = spacy_lemmatizer()
    groups = load_groups(conn)
    classify_groups(conn, groups, load_espdic_senses(args.espdic), lem)
    inflected = find_inflected_rows(conn, lem)
    write_worksheet(groups, inflected, args.out)
    verdicts: dict[str, list[int]] = defaultdict(lambda: [0, 0])
    for g in groups:
        verdicts[g.verdict][0] += 1
        verdicts[g.verdict][1] += len(g.members) - 1
    print(f"eo_word groups: {len(groups)}")
    for v, (n, extra) in sorted(verdicts.items()):
        print(f"  {v:18} {n:4} groups / {extra:4} extra concepts")
    same = [r for r in inflected if r.cid in r.lemma_cids]
    print(f"inflected EN rows: {len(inflected)} (same concept as lemma {len(same)}, "
          f"foldable {sum(foldable(r) for r in inflected)}, "
          f"cross-concept {len(inflected) - len(same)})")
    print(f"wrote {args.out}")
    return 0


def cmd_merge(conn: sqlite3.Connection, args: argparse.Namespace) -> int:
    """Merge every safe-synonym group (dry-run unless --commit)."""
    groups = load_groups(conn)
    classify_groups(conn, groups, load_espdic_senses(args.espdic), spacy_lemmatizer())
    safe = [g for g in groups if g.verdict == "safe_synonyms"]
    retire = sum(len(g.members) - 1 for g in safe)
    print(f"safe synonym groups: {len(safe)} -> {retire} concepts to retire")
    for g in safe[: args.show]:
        others = [m.cid for m in g.members if m.cid != g.survivor]
        words = [m.en_words[0][0] for m in g.members if m.en_words]
        print(f"  {g.eo_word:16} keep c{g.survivor} <- {others}  ({', '.join(words)})")
    if not args.commit:
        print("DRY-RUN — no DB changes. Re-run with --commit to merge.")
        return 0
    backup = _backup(args.lexicon, "dedup-merge")
    print(f"backup: {backup}")
    before, triples = _counts(conn), word_tier_triples(conn)
    create_concept_merge_schema(conn)
    conn.execute("BEGIN")
    try:
        moved = dropped = 0
        asof = date.today().isoformat()
        for g in safe:
            stats = merge_group(conn, g, asof)
            moved, dropped = moved + stats["moved"], dropped + stats["dropped"]
        after = _counts(conn)
        struct = structural_audit(conn)
        leftover = conn.execute(
            "SELECT COUNT(*) FROM (SELECT eo_word FROM concept WHERE eo_word IN (%s) "
            "GROUP BY eo_word HAVING COUNT(*) > 1)" % ",".join("?" * len(safe)),
            tuple(g.eo_word for g in safe),
        ).fetchone()[0] if safe else 0
        checks = {
            "word/tier triples unchanged": word_tier_triples(conn) == triples,
            "concepts retired as planned": before["concept"] - after["concept"] == retire,
            "concept_lang = before - dropped dupes":
                before["concept_lang"] - after["concept_lang"] == dropped,
            "merged eo_words now single": leftover == 0,
            **{f"{k} == 0": v == 0 for k, v in struct.items()},
        }
    except Exception as exc:  # noqa: BLE001 — any failure must roll back
        conn.execute("ROLLBACK")
        print(f"FAILED ({exc}) — rolled back, nothing written.")
        return 1
    return _finish(conn, checks, f"moved={moved} dropped_identical={dropped} "
                   f"concept {before['concept']}->{after['concept']}")


def cmd_fold(conn: sqlite3.Connection, args: argparse.Namespace) -> int:
    """Fold same-concept inflected EN rows into inflected_forms (dry-run unless --commit)."""
    rows = [r for r in find_inflected_rows(conn, spacy_lemmatizer()) if foldable(r)]
    print(f"foldable inflected rows: {len(rows)}")
    for r in rows[: args.show]:
        print(f"  {r.word:18} -> {r.lemma:16} T{r.tier} (lemma T{r.lemma_tier}) "
              f"c{r.cid} {r.source}")
    if not args.commit:
        print("DRY-RUN — no DB changes. Re-run with --commit to fold.")
        return 0
    backup = _backup(args.lexicon, "dedup-fold")
    print(f"backup: {backup}")
    before = _counts(conn)
    tiers_before = {r.word: resolved_tier(conn, r.word, r.lemma) for r in rows}
    conn.execute("BEGIN")
    try:
        inserted = sum(fold_row(conn, r) for r in rows)
        after = _counts(conn)
        changed = [
            r.word for r in rows
            if resolved_tier(conn, r.word, r.lemma) != tiers_before[r.word]
        ]
        checks = {
            "every folded word resolves to the same tier": not changed,
            "concept count unchanged": before["concept"] == after["concept"],
            "concept_lang = before - folded":
                before["concept_lang"] - after["concept_lang"] == len(rows),
            **{f"{k} == 0": v == 0 for k, v in structural_audit(conn).items()},
        }
        if changed:
            print(f"tier changes: {changed[:20]}")
    except Exception as exc:  # noqa: BLE001 — any failure must roll back
        conn.execute("ROLLBACK")
        print(f"FAILED ({exc}) — rolled back, nothing written.")
        return 1
    return _finish(conn, checks, f"folded={len(rows)} inflected_forms inserted="
                   f"{inserted} concept_lang {before['concept_lang']}->"
                   f"{after['concept_lang']}")


def _finish(conn: sqlite3.Connection, checks: dict[str, bool], summary: str) -> int:
    print(summary)
    for name, ok in checks.items():
        print(f"  AUDIT {'PASS' if ok else 'FAIL'}: {name}")
    if all(checks.values()):
        conn.commit()
        print("COMMITTED.")
        return 0
    conn.execute("ROLLBACK")
    print("AUDIT FAILED — rolled back, nothing written.")
    return 1


def main(argv: list[str] | None = None) -> int:
    """CLI entry point."""
    root = Path(__file__).resolve().parents[2]
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("command", choices=("worksheet", "merge-synonyms", "fold-inflections"))
    ap.add_argument("--lexicon", type=Path,
                    default=root / "data" / "lexicon_db" / "lexicon_v2.db")
    ap.add_argument("--espdic", type=Path,
                    default=root / "data" / "lexicon_db" / "espdic.txt")
    ap.add_argument("--out", type=Path,
                    default=root / "data" / "analysis" / "concept_dedup"
                    / "concept_dedup_review.xlsx")
    ap.add_argument("--show", type=int, default=15, help="example lines to print")
    ap.add_argument("--commit", action="store_true", help="write the DB (default: dry-run)")
    args = ap.parse_args(argv)
    if not args.lexicon.exists():
        print(f"Error: {args.lexicon} not found.", file=sys.stderr)
        return 1
    conn = sqlite3.connect(args.lexicon, isolation_level=None)
    try:
        handler = {"worksheet": cmd_worksheet, "merge-synonyms": cmd_merge,
                   "fold-inflections": cmd_fold}[args.command]
        return handler(conn, args)
    finally:
        conn.close()


if __name__ == "__main__":
    sys.exit(main())
